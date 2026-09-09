# -*- coding: utf-8 -*-
"""生成 GameReview AI 项目介绍 PPT（16:9，暗色琥珀金主题，16 页）。
用法：python docs/make_pptx.py  →  输出 docs/GameReview_AI_项目介绍.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 主题色（与全站 style.css 一致） ----------
BG    = RGBColor(0x0B, 0x0E, 0x14)   # 页面背景
SURF  = RGBColor(0x14, 0x1A, 0x25)   # 卡片
SURF2 = RGBColor(0x1A, 0x22, 0x30)   # 卡片2 / 表头
LINE  = RGBColor(0x23, 0x2C, 0x3B)   # 边框
ACC   = RGBColor(0xE8, 0xA3, 0x3D)   # 琥珀金主色
ACC2  = RGBColor(0xF2, 0xC8, 0x79)   # 浅金
TXT   = RGBColor(0xE9, 0xED, 0xF4)   # 主文本
MUT   = RGBColor(0x98, 0xA2, 0xB3)   # 次文本
DIM   = RGBColor(0x5E, 0x6A, 0x7D)   # 弱文本
BLUE  = RGBColor(0x5B, 0x8D, 0xEF)
GREEN = RGBColor(0x46, 0xC7, 0x7C)
RED   = RGBColor(0xE5, 0x69, 0x5B)

FONT = "微软雅黑"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 16


def new_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def _set_font(run, name):
    """同时设置 latin / east-asian 字体，保证中文生效。"""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def _apply(p, d):
    """按 spec 渲染一个段落。spec: {segs:[(text,bold,color)], size, align, after, before, ls, font}"""
    p.alignment = d.get("align", PP_ALIGN.LEFT)
    p.space_after = Pt(d.get("after", 4))
    p.space_before = Pt(d.get("before", 0))
    p.line_spacing = d.get("ls", 1.15)
    for (t, b, c) in d["segs"]:
        r = p.add_run()
        r.text = t
        f = r.font
        f.size = Pt(d.get("size", 12.5))
        f.bold = b
        f.color.rgb = c
        _set_font(r, d.get("font", FONT))


def P(text, size=12.5, color=MUT, bold=False, **kw):
    d = {"segs": [(text, bold, color)], "size": size}
    d.update(kw)
    return d


def B(item, size=12.5, after=5):
    """项目符号条目：item 为 str 或 [(text,bold),...]（加粗引导词用）。"""
    segs = [("▸ ", False, ACC)]
    if isinstance(item, str):
        segs.append((item, False, MUT))
    else:
        for (t, b) in item:
            segs.append((t, b, TXT if b else MUT))
    return {"segs": segs, "size": size, "after": after}


def paras_box(s, x, y, w, h, pd, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for d in pd:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _apply(p, d)
    return tb


def box(s, x, y, w, h, fill=SURF, line_c=LINE, radius=0.10):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_c is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_c
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.10)
    return shp


def fill_shape(shp, pd, anchor=MSO_ANCHOR.TOP):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for d in pd:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        _apply(p, d)
    return shp


def header(s, tag, title, subtitle, num):
    paras_box(s, 0.5, 0.28, 6, 0.3, [P(tag, size=11, color=ACC, bold=True)])
    paras_box(s, 0.5, 0.55, 12.3, 0.72, [P(title, size=30, color=ACC, bold=True)])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                             Inches(0.52), Inches(1.30), Inches(1.6), Inches(0.045))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACC
    bar.line.fill.background()
    bar.shadow.inherit = False
    paras_box(s, 0.5, 1.42, 12.3, 0.4, [P(subtitle, size=13, color=MUT)])
    paras_box(s, 12.2, 7.08, 0.9, 0.3,
              [P("%d / %d" % (num, TOTAL), size=10, color=DIM, align=PP_ALIGN.RIGHT)])


def col_title(text):
    return P(text, size=15, color=ACC, bold=True, after=10)


def pills(s, labels, y, h=0.44, size=12, gap=0.22):
    """一行居中的胶囊标签。"""
    widths = [0.42 + len(l) * 0.105 for l in labels]
    total = sum(widths) + gap * (len(labels) - 1)
    x = (13.333 - total) / 2
    for l, w in zip(labels, widths):
        shp = box(s, x, y, w, h, fill=SURF, line_c=LINE, radius=0.5)
        fill_shape(shp, [P(l, size=size, color=MUT, align=PP_ALIGN.CENTER)],
                   anchor=MSO_ANCHOR.MIDDLE)
        x += w + gap


def bullets_col(s, x, y, w, title, items, size=12.5):
    pd = [col_title(title)] + [B(it, size=size) for it in items]
    paras_box(s, x, y, w, 5.2, pd)


# ============================================================
# Slide 1 封面
# ============================================================
s = new_slide()
# 「GameReview」白 + 「AI」金，双段渲染
tb = s.shapes.add_textbox(Inches(0.5), Inches(1.9), Inches(12.33), Inches(1.1))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
for t, c in [("GameReview ", TXT), ("AI", ACC)]:
    r = p.add_run(); r.text = t
    r.font.size = Pt(52); r.font.bold = True; r.font.color.rgb = c
    _set_font(r, FONT)
paras_box(s, 0.5, 3.05, 12.33, 0.5,
          [P("Ollama 本地大模型驱动的游戏社区评测分享平台", size=18, color=MUT, align=PP_ALIGN.CENTER)])
pills(s, ["FastAPI", "Ollama qwen2.5:3b", "原生 HTML/CSS/JS", "SQLAlchemy ORM", "SQLite"], 4.05)
paras_box(s, 0.5, 4.95, 12.33, 0.4,
          [P("2026-09-07 · 毕业项目介绍演示", size=12, color=DIM, align=PP_ALIGN.CENTER)])

# ============================================================
# Slide 2 产品定位
# ============================================================
s = new_slide()
header(s, "PRODUCT", "产品定位", "以 Ollama 本地大模型为智能内核，融合评分、评测、社区与 AI 推荐", 2)
cards = [
    ("🎯", "面向玩家", "发现好游戏、参与社区互动、获取真实评测与评分参考，AI 推荐助手智能匹配兴趣"),
    ("✍️", "面向创作者", "发布评测攻略、四维评分（剧情/画面/玩法/优化）、AI 风控审核、积分等级体系"),
    ("⚙️", "面向管理员", "后台审核管理、AI 审核日志、社区举报处理、数据仪表盘实时监控"),
]
x0, cw, gap = 0.77, 3.7, 0.35
for i, (ico, t, body) in enumerate(cards):
    shp = box(s, x0 + i * (cw + gap), 2.25, cw, 2.7)
    fill_shape(shp, [
        P(ico, size=26, align=PP_ALIGN.CENTER, after=8),
        P(t, size=15, color=ACC2, bold=True, align=PP_ALIGN.CENTER, after=8),
        P(body, size=12, color=MUT, ls=1.3),
    ])

# ============================================================
# Slide 3 用户角色
# ============================================================
s = new_slide()
header(s, "USERS", "用户角色与权限", "三级角色体系 + 多账号同时登录 + 创作者等级 L0-L4", 3)
roles = [
    ("🎮", "玩家 Player", "player / player123", BLUE,
     "浏览全部内容、社区互动（发帖/评论/点赞/收藏/回复）、游戏评分与短评论、申请成为创作者"),
    ("✍️", "创作者 Creator", "creator / creator123", GREEN,
     "玩家全部权限 + 发布评测攻略（四维评分）、AI 创作助手、创作者工作台、积分等级 L0-L4"),
    ("🛡️", "管理员 Admin", "admin / admin123", RED,
     "创作者全部权限 + 后台管理（用户/游戏/评论/评测审核/社区/举报/AI 日志/仪表盘）"),
]
for i, (ico, name, badge, bc, body) in enumerate(roles):
    shp = box(s, x0 + i * (cw + gap), 2.05, cw, 2.55)
    fill_shape(shp, [
        P(ico, size=24, align=PP_ALIGN.CENTER, after=6),
        P(name, size=15, color=ACC2, bold=True, align=PP_ALIGN.CENTER, after=4),
        P(badge, size=10, color=bc, bold=True, align=PP_ALIGN.CENTER, after=8),
        P(body, size=11, color=MUT, ls=1.28),
    ])
hb = box(s, 1.17, 4.95, 11.0, 1.45, fill=SURF2, line_c=ACC)
fill_shape(hb, [
    P("🔄 多账号同时登录", size=14, color=ACC, bold=True, after=6),
    P("同一浏览器登录多个账号（localStorage gc_accounts），导航栏下拉一键切换，退出当前自动切到下一个。"
      "401 时自动移除失效账号并切换。添加账号入口：login.html?add=1", size=12, color=MUT, ls=1.3),
])

# ============================================================
# Slide 4 AI 三级降级
# ============================================================
s = new_slide()
header(s, "AI ENGINE", "AI 智能内核：三级降级", "Ollama 本地模型优先 → DeepSeek 云端备选 → 规则引擎兜底", 4)
flows = [
    ("① Ollama 本地模型（优先）",
     "模型：qwen2.5:3b · 端点：localhost:11434/api/chat\n注入：平台数据快照 + 游戏库目录 + 对话历史\n参数：temperature=0.6, num_predict=600", 1.35),
    ("② DeepSeek 云端 API（备选）",
     "模型：deepseek-chat · JSON response format\n触发条件：DEEPSEEK_API_KEY 非空", 1.05),
    ("③ 规则引擎（兜底，零 AI 也可用）",
     "意图识别 + 关键词匹配 + 游戏知识库 + 数据库查询", 0.85),
]
y = 2.0
for i, (t, body, h) in enumerate(flows):
    shp = box(s, 2.57, y, 8.2, h, fill=SURF, line_c=ACC)
    pd = [P(t, size=13.5, color=ACC, bold=True, align=PP_ALIGN.CENTER, after=4)]
    for line in body.split("\n"):
        pd.append(P(line, size=11.5, color=MUT, align=PP_ALIGN.CENTER, after=2))
    fill_shape(shp, pd, anchor=MSO_ANCHOR.MIDDLE)
    y += h
    if i < len(flows) - 1:
        paras_box(s, 2.57, y - 0.06, 8.2, 0.42,
                  [P("↓ 不可用时降级", size=12, color=ACC, bold=True, align=PP_ALIGN.CENTER)])
        y += 0.40

# ============================================================
# Slide 5 AI 对话模块详解
# ============================================================
s = new_slide()
header(s, "AI DETAILS", "AI 对话模块详解", "三层架构 · 实时数据快照 · 游戏名四层匹配", 5)
bullets_col(s, 0.7, 2.1, 5.9, "🔧 架构", [
    [("前端：", True), ("mock.js + 事件委托 data-qs/data-action + busy 防重复", False)],
    [("后端：", True), ("ai_client.py（规则 + Ollama + DeepSeek）", False)],
    [("API：", True), ("POST /api/ai/chat（匿名可用）", False)],
    "4 个 Tab：推荐游戏 / 高分游戏 / 游戏百科 / 卡关",
    [("标签轮换：", True), ('data-qs="q1|q2|q3" 竖线分隔按序轮换', False)],
    [("创作者助手：", True), ("快捷按钮 + 自由对话（需 creator 角色）", False)],
])
bullets_col(s, 6.9, 2.1, 5.9, "🧠 核心能力", [
    [("实时数据快照：", True), ("_build_platform_digest(db) 聚合统计", False)],
    "30 秒 TTL 缓存避免重复查聚合 SQL",
    "快照内容：游戏数/分类/用户/评测/社区/评分/Top5",
    [("游戏名四层匹配：", True), ("别名 → 中文全名 → 英文名 → 归一化", False)],
    [("推荐提取：", True), ("从自然语言回复提取游戏名 → 卡片", False)],
    [("强制规则：", True), ("标签专题（巡礼/百科/榜单）强制规则引擎", False)],
])

# ============================================================
# Slide 6 游戏详情页
# ============================================================
s = new_slide()
header(s, "GAME DETAIL", "游戏详情页", "三栏 Hero 布局 + 四维雷达评分 + 截图画廊", 6)
bullets_col(s, 0.7, 2.1, 5.9, "📐 三栏布局", [
    [("左栏：", True), ("Steam CDN 四图（600x900/header/capsule/hero）+ onerror SVG 兜底", False)],
    [("中栏：", True), ("meta-grid 信息、标签、描述展开/收起、Steam 购买按钮", False)],
    [("右栏：", True), ("综合评分 + 四维雷达 SVG + dim-chips + 我的评分 + 热门短评", False)],
    "下方：截图画廊（lightbox）+ 评测攻略 + 相似推荐",
])
bullets_col(s, 6.9, 2.1, 5.9, "📊 评分系统", [
    [("game_ratings 表 ", True), ("(game_id, user_id) 唯一约束", False)],
    [("score 为 int 1-10", True), ("，重复提交 = 更新", False)],
    [("综合评分 = ", True), ("评测四维均值 + 用户分数加权", False)],
    [("四维雷达仅从评测聚合", True), ("（用户评分不影响四维）", False)],
    "两个 keyword 过滤器（按游戏名 vs 按评测内容）",
])

# ============================================================
# Slide 7 评测攻略系统（表格）
# ============================================================
s = new_slide()
header(s, "REVIEWS", "评测攻略系统（PGC）", "创作者撰写 · AI 风控审核 · 积分等级体系", 7)
rows = [
    ["环节", "实现", "规则"],
    ["📝 评测结构", "标题 + 关联游戏 + 正文 + 标签 + 四维评分（剧情/画面/玩法/优化）", "仅 creator/admin 可发布"],
    ["🤖 AI 审核", "ai_audit() 检测广告引流词 + 辱骂引战词 → 风控评分", "DeepSeek → 规则引擎降级"],
    ["✅ 状态流转", "draft → audit → published / rejected / manual_review", "已通过全可见，其他仅作者可见"],
    ["✏️ 编辑权限", "L0 不可编辑；L1+ 每篇 1 次（edit_used 字段）", "积分 100/300/600/1000"],
    ["🏆 积分规则", "发布 +20 精选 +30 获赞 +0.2 收藏 +0.5", "删除扣回 + 额外 30；抄袭清零"],
    ["📋 排序", "精选 is_featured → 创作者等级 → 发布时间", "—"],
    ["📄 管理", "5 个状态分类标签 + 分页（工作台 + 个人中心）", "5 条/页 · 10 条/页"],
]
tbl_shape = s.shapes.add_table(len(rows), 3, Inches(0.87), Inches(2.05),
                               Inches(11.6), Inches(4.6))
tbl = tbl_shape.table
for i, wd in enumerate([1.7, 6.1, 3.8]):
    tbl.columns[i].width = Inches(wd)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(0.42 if r == 0 else 0.58)
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURF2 if r == 0 else SURF
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.09)
        cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        run.font.size = Pt(12) if r == 0 else Pt(11.5)
        run.font.bold = (r == 0)
        run.font.color.rgb = ACC if r == 0 else MUT
        _set_font(run, FONT)

# ============================================================
# Slide 8 社区系统
# ============================================================
s = new_slide()
header(s, "COMMUNITY", "社区系统（UGC）", "帖子 · 图片上传 · 楼中楼评论 · 频率限制", 8)
bullets_col(s, 0.7, 2.1, 5.9, "📢 帖子功能", [
    "标题（2-100 字）+ 正文 + 图片（≤9 张，≤5MB/张）",
    "标签单选：游戏 / 闲聊 / 攻略 / 吐槽",
    "排序：推荐（热度）/ 最新 / 标签筛选",
    "频率限制：发帖 30s/次，评论 10s/次",
    [("图片上传：", True), ("FormData → /api/community/upload → static/uploads", False)],
])
bullets_col(s, 6.9, 2.1, 5.9, "💬 互动功能", [
    "点赞、收藏、评论（楼中楼回复 parent_id）",
    "删除自己的帖子 / 评论",
    "我的收藏/点赞仅本人可见（他人 403）",
    "帖主可删除其下所有评论",
    "管理员可删除违规帖 + 封禁用户",
])

# ============================================================
# Slide 9 评论系统
# ============================================================
s = new_slide()
header(s, "COMMENTS", "评论系统", "统一模型 · 楼中楼回复 · 删除权限", 9)
bullets_col(s, 0.7, 2.1, 5.9, "💬 统一评论模型", [
    [("Comment 表，target_type 区分：", True), ("game（短评）/ review（评测评论）", False)],
    "parent_id 支持楼中楼嵌套回复",
    [("DELETE /comments/{id}：", True), ("仅本人或 admin 可删", False)],
    [("can_delete 字段：", True), ("前端据此渲染删除按钮", False)],
    "级联删除子回复",
])
bullets_col(s, 6.9, 2.1, 5.9, "🎮 游戏短评特色", [
    [("自动精选：", True), ("最高赞 + 最早创建（后端运行时计算）", False)],
    "分页：每页 5-10 条（page_size clamp 1-10）",
    "显示顺序：精选 → 普通分页（时间倒序）",
    [("「我的评论」：", True), ("GET /comments/mine 聚合两种类型", False)],
    [("点击跳转 ", True), ("#comment-{id} 锚点直达原文", False)],
])

# ============================================================
# Slide 10 消息通知
# ============================================================
s = new_slide()
header(s, "NOTIFICATIONS", "消息通知系统", "10 秒轮询 · 5 类通知 · 直达操作页 · 全埋点", 10)
bullets_col(s, 0.7, 2.1, 5.9, "🔔 通知类型与埋点", [
    [("评论回复：", True), ("有人回复你的评论 → notify.push", False)],
    [("点赞：", True), ("评测/评论被赞 → 通知作者（自己不通知自己）", False)],
    [("收藏：", True), ("评测被收藏 → 通知作者", False)],
    [("审核结果：", True), ("通过/驳回 → 通知创作者", False)],
    [("创作者申请：", True), ("提交→通知管理员；通过/驳回→通知申请人", False)],
])
bullets_col(s, 6.9, 2.1, 5.9, "⚡ 实时性设计", [
    "导航栏铃铛 10s 轮询未读数",
    "消息中心 10s 轮询通知列表",
    "5 个 Tab + 只看未读 + 全部已读 + 分页",
    [("点击消息 ", True), ("→ 标记已读 + 直达对应页", False)],
    [("notify.push_to_admins", True), (" 通知全部管理员", False)],
])

# ============================================================
# Slide 11 管理后台（表格）
# ============================================================
s = new_slide()
header(s, "ADMIN", "管理后台", "7 大模块 · 5-8 秒轮询 · 分页管理", 11)
rows = [
    ["模块", "功能", "特色"],
    ["📊 数据仪表盘", "平台概览统计", "实时数据"],
    ["👥 用户管理", "用户列表 + 创作者申请审批", "5-8s 轮询，离开销毁定时器"],
    ["⚖️ 审核队列", "待审核评测 + 通过/驳回 + AI 日志", "5-8s 轮询，离开销毁定时器"],
    ["🎮 游戏管理", "游戏 CRUD + 分类管理", "10 条/页分页"],
    ["💬 评论管理", "全站评论列表 + 删除违规", "10 条/页分页"],
    ["📢 社区管理", "举报队列 + 帖子管理", "关键词搜索 + 查看跳转"],
    ["📋 已发布评测", "已通过评测列表", "分页浏览"],
]
tbl_shape = s.shapes.add_table(len(rows), 3, Inches(0.87), Inches(2.05),
                               Inches(11.6), Inches(4.6))
tbl = tbl_shape.table
for i, wd in enumerate([2.6, 5.4, 3.6]):
    tbl.columns[i].width = Inches(wd)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(0.44 if r == 0 else 0.58)
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = SURF2 if r == 0 else SURF
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.09)
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        run.font.size = Pt(12) if r == 0 else Pt(11.5)
        run.font.bold = (r == 0)
        run.font.color.rgb = ACC if r == 0 else MUT
        _set_font(run, FONT)

# ============================================================
# Slide 12 个人中心 & 创作者工作台
# ============================================================
s = new_slide()
header(s, "PROFILE", "个人中心 & 创作者工作台", "锚点导航 · 状态分类 · 分页管理", 12)
bullets_col(s, 0.7, 2.1, 5.9, "👤 个人中心", [
    "锚点快捷导航（收藏/评测/评论/社区/游玩档案）",
    "我的评测：5 个状态标签 + 分页（5 条/页）",
    "我的评论：聚合两种评论，点击跳转原文 + 删除",
    "评测互动：我赞过的 / 我收藏的",
    "社区：我的帖子 / 收藏 / 点赞（三 Tab）",
    "创作者申请：理由 + 擅长 + 经历",
])
bullets_col(s, 6.9, 2.1, 5.9, "✍️ 创作者工作台", [
    "评测管理表格：5 个状态标签 + 分页（10 条/页）",
    "操作：查看 / 编辑 / 提交审核 / 删除",
    "AI 攻略创作助手（快捷按钮 + 自由对话）",
    "创作者统计数据",
    "轮询防抖签名纳入 page+status",
])

# ============================================================
# Slide 13 智能返回 & 多账号
# ============================================================
s = new_slide()
header(s, "UX", "智能返回 & 多账号", "三级返回策略 + 无缝账号切换", 13)
bullets_col(s, 0.7, 2.1, 5.9, "↩️ 智能返回 U.goBack()", [
    [("同标签页", True), ("（history>1 且非弹出）→ history.back()", False)],
    [("新标签页", True), ("（target=_blank，如管理员查看）→ document.referrer", False)],
    [("兜底", True), ("→ fallback URL（如 game-list.html）", False)],
    "覆盖 6 个页面：game-detail / review-detail / review-list / community-post / community-new / creator-edit",
])
bullets_col(s, 6.9, 2.1, 5.9, "🔐 多账号登录", [
    [("gc_accounts", True), (" 存储已登录账号列表 [{token, user}]", False)],
    [("setSession", True), (" 自动去重（按 user.id）", False)],
    "导航栏下拉展示全部账号，当前 ✓ 标记",
    [("一键切换（", True), ("switchAccount → reload）", False)],
    "退出当前自动切到下一个",
    "401 自动移除失效账号 + 切换",
])

# ============================================================
# Slide 14 技术架构
# ============================================================
s = new_slide()
header(s, "ARCHITECTURE", "技术架构", "FastAPI + Ollama 本地模型 + 原生前端 · 单进程按需启动", 14)
bullets_col(s, 0.7, 2.1, 5.9, "🔧 技术栈", [
    [("后端：", True), ("FastAPI ≥0.110 + SQLAlchemy ≥1.4 + SQLite", False)],
    [("AI：", True), ("Ollama qwen2.5:3b（本地）+ DeepSeek（备选）+ 规则引擎", False)],
    [("前端：", True), ("原生 HTML/CSS/JS（无框架）+ 模块化", False)],
    [("认证：", True), ("python-jose JWT + bcrypt", False)],
    [("部署：", True), ("uvicorn 单进程 + start-platform.bat 一键启动", False)],
    [("HTTP：", True), ("httpx 异步客户端", False)],
])
tree = """app/
  routers/       12 个路由文件
  models.py · schemas.py
  services.py · repositories.py
  ai_client.py · notify.py
  core/  config.py · deps.py · points.py
  db/    base.py · init_db.py
static/
  js/    api / common / auth / mock
  css/   style.css (proto12)
  18 个前台 HTML
  admin/ 6 个后台 HTML
docs/    PRD_GameReview_AI.md
main.py · .env · requirements.txt"""
shp = box(s, 6.9, 2.05, 5.7, 4.6)
fill_shape(shp, [P(tree, size=10.5, color=MUT, font=MONO, ls=1.25)])

# ============================================================
# Slide 15 种子数据与统计
# ============================================================
s = new_slide()
header(s, "DATA", "种子数据与统计", "完整 Mock 数据覆盖全功能场景", 15)
stats = [("7", "用户"), ("30", "游戏"), ("7", "分类"), ("67", "评测"),
         ("12", "API 路由"), ("24", "页面")]
x0, w, gap = 0.87, 1.6, 0.4
for i, (num, label) in enumerate(stats):
    shp = box(s, x0 + i * (w + gap), 2.2, w, 1.5)
    fill_shape(shp, [
        P(num, size=38, color=ACC, bold=True, align=PP_ALIGN.CENTER, after=2),
        P(label, size=12, color=MUT, align=PP_ALIGN.CENTER),
    ], anchor=MSO_ANCHOR.MIDDLE)
hb = box(s, 1.4, 4.35, 10.53, 2.0, fill=SURF2, line_c=ACC)
fill_shape(hb, [
    P("🔑 关键设计决策", size=14, color=ACC, bold=True, after=8),
    P("① AI 三级降级（Ollama→DeepSeek→规则）保证零 AI 也可用   ② 实时数据快照 30s TTL 缓存   "
      "③ notify.push 自己不通知自己   ④ CSS 变量重映射保留旧类名兼容", size=12, color=MUT, ls=1.4, after=6),
    P("⑤ 按需一键启动非开机自启   ⑥ Comment.target_type 统一游戏短评/评测评论   "
      "⑦ can_delete 字段前端渲染删除按钮   ⑧ Ollama 不参与审核，审核仅 DeepSeek/规则引擎",
      size=12, color=MUT, ls=1.4),
])

# ============================================================
# Slide 16 结尾
# ============================================================
s = new_slide()
tb = s.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.33), Inches(1.1))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
for t, c in [("谢谢 ", TXT), ("观看", ACC)]:
    r = p.add_run(); r.text = t
    r.font.size = Pt(52); r.font.bold = True; r.font.color.rgb = c
    _set_font(r, FONT)
paras_box(s, 0.5, 3.15, 12.33, 0.5,
          [P("GameReview AI — Ollama 本地大模型驱动的游戏社区评测分享平台",
             size=17, color=MUT, align=PP_ALIGN.CENTER)])
pills(s, ["🤖 Ollama qwen2.5:3b", "🎮 30 款游戏", "✍️ PGC 评测",
          "📢 UGC 社区", "🔔 实时通知", "🔐 多账号"], 4.15)
paras_box(s, 0.5, 5.1, 12.33, 0.8, [
    P("访问地址：http://127.0.0.1:8000/static/index.html", size=12, color=DIM, align=PP_ALIGN.CENTER, after=4),
    P("PRD 文档：docs/PRD_GameReview_AI.md", size=12, color=DIM, align=PP_ALIGN.CENTER),
])

# ---------- 保存 ----------
out = Path(__file__).resolve().parent / "GameReview_AI_项目介绍.pptx"
prs.save(str(out))
print("saved:", out)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
